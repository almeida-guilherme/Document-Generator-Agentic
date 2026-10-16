"""
Writes a process into a company template: the template's own file is the starting point, so its cover,
styles, headers, numbering and fixed text stay exactly as the company has them. Blocks are removed,
rewritten or filled as the template analysis says (see template.py), the detailed steps and the flowchart
go where the template's examples were, and the table of contents is rebuilt from the headings.
"""
import copy
import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches
from docx.text.paragraph import Paragraph
from PIL import Image

from flowchart import build_flowchart
from schemas import ProcessModel
from template import TemplateAnalysis, TemplateFill, body_blocks, is_toc

MAX_IMAGE_WIDTH = Inches(6.3)
MAX_IMAGE_HEIGHT = Inches(8)
# The flowchart shares its page with the section heading, description and caption.
MAX_DIAGRAM_HEIGHT = Inches(6.3)


# --- Text helpers that keep the template's formatting ---

def _runs_with_text(p_el) -> list:
    return [r for r in p_el.iter(qn("w:r")) if r.find(qn("w:t")) is not None]


def _plain_rpr(run_el):
    """A copy of a run's formatting without the colour or highlight that marks template placeholders."""
    rpr = run_el.find(qn("w:rPr"))
    rpr = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
    for tag in ("w:color", "w:highlight", "w:shd"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    return rpr


def _new_run(text: str, rpr=None, missing: bool = False):
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    if missing:
        rpr_el = run.find(qn("w:rPr"))
        if rpr_el is None:
            rpr_el = OxmlElement("w:rPr")
            run.insert(0, rpr_el)
        highlight = OxmlElement("w:highlight")
        highlight.set(qn("w:val"), "yellow")
        rpr_el.append(highlight)
    for n, line in enumerate(text.split("\n")):
        if n:
            run.append(OxmlElement("w:br"))
        for k, piece in enumerate(line.split("\t")):
            if k:
                run.append(OxmlElement("w:tab"))
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = piece
            run.append(t)
    return run


def set_paragraph_text(p_el, text: str, missing: bool = False, keep_color: bool = False) -> None:
    """Replaces a paragraph's runs by one run with `text`, in the formatting of its first run."""
    runs = _runs_with_text(p_el)
    rpr = None
    if runs:
        rpr = copy.deepcopy(runs[0].find(qn("w:rPr"))) if keep_color else _plain_rpr(runs[0])
    for child in list(p_el):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:ins"), qn("w:del"), qn("w:smartTag")):
            p_el.remove(child)
    p_el.append(_new_run(text, rpr, missing))


def set_cell_text(tc, text: str | None, missing_label: str) -> None:
    paragraphs = tc.findall(qn("w:p"))
    for extra in paragraphs[1:]:
        tc.remove(extra)
    set_paragraph_text(paragraphs[0], missing_label if text is None else text, missing=text is None)


def replace_in_paragraph(p_el, old: str, new: str | None, missing_label: str, ignore_case: bool = True) -> bool:
    """
    Replaces `old` in a paragraph even when Word split it over several runs, keeping the formatting of the
    run it starts in. With `new` None, the placeholder stays and is highlighted.
    """
    found, search_from = False, 0
    while True:
        runs = _runs_with_text(p_el)
        texts = ["".join(t.text or "" for t in r.findall(qn("w:t"))) for r in runs]
        full = "".join(texts)
        start = (full.lower().find(old.lower(), search_from) if ignore_case else full.find(old, search_from))
        if start < 0 or (new is None and found):
            return found
        found = True
        end = start + len(old)
        search_from = start + len(new or "")
        pos = 0
        for run, text in zip(runs, texts):
            r_start, r_end = pos, pos + len(text)
            pos = r_end
            if r_end <= start or r_start >= end:
                continue
            before = text[:max(0, start - r_start)]
            after = text[max(0, end - r_start):] if end < r_end else ""
            if new is None:
                # Leave the text, highlight the runs it covers so the author sees what's missing.
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    run.insert(0, rpr)
                if rpr.find(qn("w:highlight")) is None:
                    highlight = OxmlElement("w:highlight")
                    highlight.set(qn("w:val"), "yellow")
                    rpr.append(highlight)
                continue
            first = r_start <= start
            ts = run.findall(qn("w:t"))
            for t in ts[1:]:
                run.remove(t)
            ts[0].set(qn("xml:space"), "preserve")
            ts[0].text = before + (new if first else "") + after
            if first:
                # Placeholders are usually coloured to stand out; the value takes the text's normal colour.
                rpr = run.find(qn("w:rPr"))
                if rpr is not None and not before and not after:
                    for el in rpr.findall(qn("w:color")):
                        rpr.remove(el)
        if new is None:
            return found


def _all_paragraphs(doc) -> list:
    """Paragraphs of the body (tables included) and of every page header and footer."""
    paragraphs = list(doc.element.body.iter(qn("w:p")))
    seen = set()
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part.is_linked_to_previous or id(part.part) in seen:
                continue
            seen.add(id(part.part))
            paragraphs += part._element.iter(qn("w:p"))
    return paragraphs


# --- Blocks ---

def _picture_paragraph(doc, image_path: str, template_p=None, max_height=MAX_IMAGE_HEIGHT):
    """A centred paragraph holding the image, scaled to fit the page; returned detached for placing."""
    usable = MAX_IMAGE_WIDTH
    section = doc.sections[0]
    if section.page_width and section.left_margin is not None and section.right_margin is not None:
        usable = min(usable, section.page_width - section.left_margin - section.right_margin)
    with Image.open(image_path) as img:
        width, height = img.size
    scale = min(usable / width, max_height / height)
    p = doc.add_paragraph()
    if template_p is not None and template_p.find(qn("w:pPr")) is not None:
        p._p.insert(0, copy.deepcopy(template_p.find(qn("w:pPr"))))
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(image_path, width=Emu(int(width * scale)), height=Emu(int(height * scale)))
    p.paragraph_format.keep_with_next = False
    el = p._p
    el.getparent().remove(el)
    return el


def _clone_paragraph(p_el, text: str, missing: bool = False):
    new = copy.deepcopy(p_el)
    set_paragraph_text(new, text, missing)
    return new


def fill_table(tbl, plan, fill, missing_label: str) -> None:
    rows = tbl.findall(qn("w:tr"))
    header, data = rows[:plan.header_rows], rows[plan.header_rows:]
    if plan.action == "table_answers":
        for tr, answer in zip(data, fill.answers if fill else []):
            cells = tr.findall(qn("w:tc"))
            if plan.answer_column < len(cells):
                set_cell_text(cells[plan.answer_column], answer, missing_label)
        for tr in data[len(fill.answers) if fill else 0:]:
            cells = tr.findall(qn("w:tc"))
            if plan.answer_column < len(cells):
                set_cell_text(cells[plan.answer_column], None, missing_label)
        return
    prototype = copy.deepcopy(data[0] if data else header[-1])
    for tr in data:
        tbl.remove(tr)
    values = (fill.rows if fill else []) or [[None]]
    for row in values:
        tr = copy.deepcopy(prototype)
        cells = tr.findall(qn("w:tc"))
        row = list(row) + [None] * (len(cells) - len(row))
        # A row the model left completely empty gets one marker, in its widest cell (not in a narrow "#").
        marked = _widest(cells) if all(v is None for v in row) else None
        for tc, value in zip(cells, row):
            if value is None and marked is not None and tc is not marked:
                value = ""
            set_cell_text(tc, value, missing_label)
        tbl.append(tr)


def _widest(cells: list):
    def width(tc):
        w = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
        return int(w.get(qn("w:w"), 0)) if w is not None and (w.get(qn("w:w")) or "").isdigit() else 0
    return max(cells, key=width)


def _step_paragraph(example_p, label: str, separator: str, text: str):
    """A step written like the template's example step: its label run's format, then its text run's format."""
    new = copy.deepcopy(example_p)
    runs = _runs_with_text(new)
    label_rpr = runs[0].find(qn("w:rPr")) if runs else None
    text_rpr = _plain_rpr(runs[-1]) if runs else None
    for child in list(new):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:ins"), qn("w:del"), qn("w:smartTag")):
            new.remove(child)
    new.append(_new_run(label, label_rpr))
    new.append(_new_run(separator + text, text_rpr))
    return new


def _heading_paragraph(example_heading, text: str):
    new = copy.deepcopy(example_heading)
    for el in new.findall(qn("w:bookmarkStart")) + new.findall(qn("w:bookmarkEnd")):
        new.remove(el)
    set_paragraph_text(new, text, keep_color=True)
    return new


def build_steps(doc, model: ProcessModel, analysis: TemplateAnalysis, fill: TemplateFill, blocks: list,
                stats: dict) -> list:
    """The detailed steps as paragraphs (group headings, steps, decisions, screenshots), not yet placed."""
    fmt = analysis.step_format
    example_p = blocks[fmt.example_step_index] if fmt.example_step_index is not None else None
    if example_p is None:
        example_p = doc.add_paragraph()._p
        example_p.getparent().remove(example_p)
    heading = blocks[fmt.example_group_heading_index] if fmt.example_group_heading_index is not None else None
    starts = {g.first_step: g.title for g in fill.step_groups} if heading is not None else {}

    # The model sometimes copies the template's numbering into the text ("P4 - Abrir..."); the label is added here.
    own_label = re.compile(r"^\s*" + re.escape(fmt.label).replace(re.escape("{n}"), r"\d+") + r"\s*[-–:.]?\s*")
    out = []
    for step in model.as_is:
        if step.number in starts:
            out.append(_heading_paragraph(heading, starts[step.number]))
        p = _step_paragraph(example_p, fmt.label.replace("{n}", str(step.number)), fmt.separator,
                            own_label.sub("", step.instruction, count=1))
        out.append(p)
        if step.decision:
            out.append(_step_paragraph(example_p, f"{analysis.labels.decision}:", " ", step.decision))
        for text_p in out[-2 if step.decision else -1:]:
            # A step's text stays on the page of its screenshot.
            Paragraph(text_p, None).paragraph_format.keep_with_next = True
        image = step.main_frame_path
        if image and os.path.exists(image):
            out.append(_picture_paragraph(doc, image))
            stats["images_inserted"] += 1
        else:
            stats["images_missing"].append(step.number)
    return out


# --- Table of contents ---

def _heading_level(doc, p_el) -> int | None:
    name = Paragraph(p_el, doc).style.name.lower() if Paragraph(p_el, doc).style is not None else ""
    if name.startswith("heading") and name.split()[-1].isdigit():
        return int(name.split()[-1])
    return None


def rebuild_toc(doc, old_numbers: dict[str, str]) -> None:
    """
    Rewrites the cached entries of the table of contents field from the document's headings, so it lists
    the new step groups instead of the template's examples. Word recomputes it (with page numbers) when the
    document is opened, since updateFields is set.
    """
    body = doc.element.body
    toc = [p for p in body.iter(qn("w:p")) if is_toc(Paragraph(p, doc))]
    if not toc:
        return
    first = toc[0]
    # The field's start (begin, instruction, separate) is in the first entry; keep it, drop the entries.
    field_start = []
    for child in first:
        if child.tag == qn("w:r"):
            field_start.append(copy.deepcopy(child))
            sep = child.find(qn("w:fldChar"))
            if sep is not None and sep.get(qn("w:fldCharType")) == "separate":
                break
    styles_by_level = {}
    for p in toc:
        style = p.find(qn("w:pPr"))
        level = re.search(r"\d", Paragraph(p, doc).style.name or "")
        if style is not None and level:
            styles_by_level.setdefault(int(level.group()), style)
    end_p = toc[-1]

    entries = []
    counters: dict[int, int] = {}
    numbers: dict[int, str] = {}
    bookmark_id = 9000
    for p in body.iter(qn("w:p")):
        level = _heading_level(doc, p)
        if level is None or level > 3 or p in toc:
            continue
        text = Paragraph(p, doc).text.strip()
        if not text:
            continue
        number = old_numbers.get(text)
        if number is None:
            counters[level] = counters.get(level, 0) + 1
            number = f"{numbers.get(level - 1, '')}{counters[level]}." if level > 1 else ""
        else:
            counters[level] = int(re.findall(r"\d+", number)[-1]) if re.findall(r"\d+", number) else counters.get(level, 0)
        numbers[level] = number
        for deeper in [lv for lv in counters if lv > level]:
            counters[deeper] = 0
        name = f"_TocPdd{bookmark_id}"
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        ppr = p.find(qn("w:pPr"))
        (ppr.addnext(start) if ppr is not None else p.insert(0, start))
        p.append(end)
        bookmark_id += 1

        entry = OxmlElement("w:p")
        style = styles_by_level.get(level)
        if style is None and styles_by_level:
            style = styles_by_level[min(styles_by_level)]
        if style is not None:
            entry.append(copy.deepcopy(style))
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), name)
        link.append(_new_run(f"{number}\t{text}" if number else text))
        entry.append(link)
        entries.append(entry)

    if not entries:
        return
    position = 1 if entries[0].find(qn("w:pPr")) is not None else 0
    for offset, child in enumerate(field_start):
        entries[0].insert(position + offset, child)
    # Keep the paragraph holding the field's end (the template's last entry), emptied of its old entry.
    for p in toc:
        if p is not end_p:
            p.getparent().remove(p)
    for child in list(end_p):
        if child.tag == qn("w:hyperlink"):
            end_p.remove(child)
    for entry in entries:
        end_p.addprevious(entry)


def _old_toc_numbers(doc) -> dict[str, str]:
    """Heading text -> number shown in the template's table of contents ("4.1."), which Word computed."""
    numbers = {}
    for p in doc.element.body.iter(qn("w:p")):
        if is_toc(Paragraph(p, doc)):
            text = Paragraph(p, doc).text
            parts = [s.strip() for s in text.split("\t")]
            if len(parts) >= 3 and re.fullmatch(r"[\d.]+", parts[0]):
                numbers[parts[1]] = parts[0]
    return numbers


# --- Clean-up ---

def _fix_bookmarks(doc) -> None:
    """Removing blocks can leave a bookmark's start without its end (or the reverse); Word dislikes both."""
    body = doc.element.body
    starts = {b.get(qn("w:id")): b for b in body.iter(qn("w:bookmarkStart"))}
    ends = {b.get(qn("w:id")): b for b in body.iter(qn("w:bookmarkEnd"))}
    for bid in set(starts) ^ set(ends):
        el = starts.get(bid) if bid in starts else ends[bid]
        el.getparent().remove(el)


def _update_fields_on_open(doc) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)


def render_template(template_path: str, analysis: TemplateAnalysis, fill: TemplateFill, model: ProcessModel,
                    output_path: str) -> tuple[str, dict]:
    doc = Document(template_path)
    blocks = body_blocks(doc)
    stats = {"images_inserted": 0, "images_missing": [], "flowchart": False, "missing_marks": 0,
             "template": os.path.basename(template_path)}
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    missing = analysis.missing_label
    old_numbers = _old_toc_numbers(doc)
    fills = {f.index: f for f in fill.blocks}
    remove = []

    for plan in analysis.blocks:
        el = blocks[plan.index]
        if plan.action == "remove":
            remove.append(el)
        elif plan.action == "write":
            paragraphs = (fills[plan.index].paragraphs if plan.index in fills else []) or [None]
            for text in paragraphs:
                el.addprevious(_clone_paragraph(el, missing if text is None else text, missing=text is None))
            remove.append(el)
        elif plan.action in ("table_rows", "table_answers"):
            fill_table(el, plan, fills.get(plan.index), missing)
        elif plan.action == "diagram":
            try:
                chart = build_flowchart(model, os.path.join(os.path.dirname(output_path), "flowchart"),
                                        start_label=analysis.labels.start, end_label=analysis.labels.end)
                el.addprevious(_picture_paragraph(doc, chart, el, MAX_DIAGRAM_HEIGHT))
                stats["flowchart"] = True
            except Exception as e:
                print(f"Failed to generate/insert flowchart: {e}")
                el.addprevious(_clone_paragraph(el, missing, missing=True))
            remove.append(el)

    # The steps go where the template's example steps start (their group heading, if any).
    fmt = analysis.step_format
    anchors = [i for i in (fmt.example_group_heading_index, fmt.example_step_index) if i is not None]
    steps = build_steps(doc, model, analysis, fill, blocks, stats)
    if anchors:
        anchor = blocks[min(anchors)]
        for p in steps:
            anchor.addprevious(p)
        if anchor not in remove:
            remove.append(anchor)
    else:
        sect_pr = doc.element.body.find(qn("w:sectPr"))
        for p in steps:
            sect_pr.addprevious(p) if sect_pr is not None else doc.element.body.append(p)

    for el in remove:
        if el.getparent() is not None:
            el.getparent().remove(el)

    # Exact spellings first, so "<Nome do processo>" and "<nome do processo>" each get their own value;
    # then other capitalizations of the same placeholder (e.g. "<cliente>" when only "<Cliente>" was listed).
    values = {p.text: p.value for p in fill.placeholders}
    paragraphs = _all_paragraphs(doc)
    for ignore_case in (False, True):
        for placeholder in analysis.placeholders:
            for p_el in paragraphs:
                replace_in_paragraph(p_el, placeholder.text, values.get(placeholder.text), missing, ignore_case)

    rebuild_toc(doc, old_numbers)
    _fix_bookmarks(doc)
    _update_fields_on_open(doc)
    stats["missing_marks"] = sum(1 for h in doc.element.body.iter(qn("w:highlight")) if h.get(qn("w:val")) == "yellow")
    doc.save(output_path)
    return output_path, stats
