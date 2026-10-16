"""Company templates: analysis decisions applied to the template's own file, offline with fake model answers."""
import json
import re
import pytest
from docx import Document
from docx.oxml.ns import qn

from pipeline import run_pdd_pipeline
from render_template import replace_in_paragraph
from schemas import repair_text
from template import (BlockFill, BlockPlan, Labels, Placeholder, PlaceholderValue, StepFormat, StepGroup,
                      TemplateAnalysis, TemplateFill, normalize_fill, body_blocks, describe_template)
from tests.synthetic_video import build
from tests.template_fixture import build_template
from tests.test_pipeline import FakeLLM


def analysis_for_fixture() -> TemplateAnalysis:
    plan = {2: "table_rows", 6: "remove", 7: "table_rows", 9: "table_answers", 11: "diagram", 12: "write",
            14: "remove", 15: "remove", 16: "remove"}
    return TemplateAnalysis(
        language="Brazilian Portuguese", missing_label="[A preencher]",
        placeholders=[Placeholder(text="<Cliente>", meaning="client"),
                      Placeholder(text="<nome do processo>", meaning="process name"),
                      Placeholder(text="<Nome do processo>", meaning="process name"),
                      Placeholder(text="<Lista de todos os sistemas envolvidos no processo>", meaning="instruction")],
        blocks=[BlockPlan(index=i, action=plan.get(i, "keep"), purpose="x") for i in range(17)],
        step_format=StepFormat(example_step_index=15, example_group_heading_index=14, label="P{n}", separator=" - "),
        labels=Labels(decision="Decisão", start="Início", end="Fim"))


class FakeTemplateLLM(FakeLLM):
    """FakeLLM that also answers the template analysis and fill."""

    def __init__(self):
        super().__init__()
        self.fill_prompt = ""
        self.system_prompts: list[str] = []

    def _answer(self, schema, messages):
        self.system_prompts.append(messages[0].content)
        if schema is TemplateAnalysis:
            return analysis_for_fixture()
        if schema is TemplateFill:
            self.fill_prompt = messages[-1].content
            return TemplateFill(
                placeholders=[PlaceholderValue(text="<Cliente>", value="?"),
                              PlaceholderValue(text="<nome do processo>", value="emissão de faturas"),
                              PlaceholderValue(text="<Nome do processo>", value="Emissão de faturas")],
                blocks=[BlockFill(index=2, rows=[["0.1", "16/10/2026", "null"]]),
                        BlockFill(index=7, rows=[["1", "Intranet", "?"]]),
                        BlockFill(index=9, answers=["2 pessoas", "?"]),
                        BlockFill(index=12, paragraphs=["Figura 1 - Diagrama do processo"])],
                step_groups=[StepGroup(title="Preparar", first_step=1), StepGroup(title="Emitir", first_step=4),
                             StepGroup(title="Fora do processo", first_step=99)])
        return super()._answer(schema, messages)


@pytest.fixture(scope="module")
def video(tmp_path_factory):
    path = tmp_path_factory.mktemp("video") / "synthetic.mp4"
    build(str(path))
    return path


@pytest.fixture
def template(tmp_path):
    return build_template(str(tmp_path / "PDD_Template.docx"), str(tmp_path / "example.png"))


def test_fixture_indexes_match(template):
    """The fake analysis refers to blocks by index; this keeps the fixture and those indexes in step."""
    text = describe_template(template)
    assert "[11] paragraph style='Normal' [IMAGE]" in text and "[15] paragraph" in text and "P1 - Realizar" in text
    assert "Page headers/footers: <cliente> <Nome do processo>" in text


def _texts(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_document_is_written_into_the_template(video, template, tmp_path):
    llm = FakeTemplateLLM()
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    output, report = run_pdd_pipeline(str(video), str(run_dir), llm=llm, structure_llm=llm,
                                      transcribe=lambda v, audio_dir: [], template_path=template)

    for name in ("template.docx", "template_analysis.json", "template_fill.json"):
        assert (run_dir / name).exists()
    # Steps are written in the template's language and in the style of its example step.
    assert any("Brazilian Portuguese" in p and "Realizar abertura do SAP 760" in p for p in llm.system_prompts)
    assert "[9] table_answers" in llm.fill_prompt and "Quantas pessoas executam?" in llm.fill_prompt

    doc = Document(output)
    text = _texts(doc)
    # The example from another process and the instructions are gone.
    for leftover in ("Login SAP", "Realizar abertura", "Outlook 2011", "Maria Exemplo", "<Lista de todos",
                     "Exemplo de diagrama"):
        assert leftover not in text
    # Fixed text kept, placeholders filled (also when split over runs, and in the page header).
    assert "Este documento descreve o processo emissão de faturas escolhido para automação." in text
    assert doc.sections[0].header.paragraphs[0].text == "<cliente> Emissão de faturas"
    # Steps in the template's format, grouped under its kind of heading; the out-of-range group is ignored.
    steps = [p for p in doc.paragraphs if re.match(r"P\d+ - ", p.text)]
    assert len(steps) == report["steps"]
    assert steps[0].runs[0].text == "P1" and steps[0].runs[0].bold
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["Preparar", "Emitir"]
    assert any(p.text.startswith("Decisão: If the client exists") for p in doc.paragraphs)
    # Tables: example rows replaced, answers filled, unknown values highlighted for the author.
    revision, systems, metrics = doc.tables
    assert [c.text for c in revision.rows[1].cells] == ["0.1", "16/10/2026", "[A preencher]"]
    assert len(revision.rows) == 2
    assert [c.text for c in systems.rows[1].cells] == ["1", "Intranet", "[A preencher]"]
    assert [r.cells[1].text for r in metrics.rows] == ["Descrição", "2 pessoas", "[A preencher]"]
    highlighted = [r for r in doc.element.body.iter(qn("w:highlight"))]
    assert report["missing_marks"] == len(highlighted) >= 4
    assert "Figura 1 - Diagrama do processo" in text
    assert report["images_inserted"] + len(report["images_missing"]) == report["steps"]
    # The template's own example image was replaced (diagram) or removed (step screenshot).
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_resumed_run_keeps_its_template(video, template, tmp_path):
    llm = FakeTemplateLLM()
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    run_pdd_pipeline(str(video), str(run_dir), llm=llm, structure_llm=llm, transcribe=lambda v, audio_dir: [],
                     template_path=template)
    (run_dir / "output" / "PDD_final.docx").unlink()
    again = FakeTemplateLLM()
    output, report = run_pdd_pipeline(str(video), str(run_dir), llm=again, structure_llm=again,
                                      transcribe=lambda v, audio_dir: [])
    assert report["template"] == "template.docx" and again.system_prompts == []  # nothing asked again
    assert "Emissão de faturas" in Document(output).sections[0].header.paragraphs[0].text


def test_placeholder_split_over_runs(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Cliente: ")
    for part in ("<Nom", "e do pro", "cesso>"):
        p.add_run(part).bold = True
    p.add_run(" fim")
    assert replace_in_paragraph(p._p, "<nome do processo>", "Faturas", "[?]")
    assert p.text == "Cliente: Faturas fim"
    assert p.runs[1].bold  # the value takes the formatting of the run the placeholder started in
    q = doc.add_paragraph("<Cliente> e <cliente>")
    assert replace_in_paragraph(q._p, "<Cliente>", None, "[?]")  # unknown: kept and highlighted
    assert q.text == "<Cliente> e <cliente>"
    assert q._p.find(".//" + qn("w:highlight")) is not None


def test_garbled_accents_are_repaired():
    assert repair_text("Abrir exporta\x00e7\x00e3o SAP") == "Abrir exportação SAP"
    assert repair_text("Ok\x07") == "Ok"


def test_unknown_answers_and_lists_arenormalize_filld():
    fill = TemplateFill(placeholders=[PlaceholderValue(text="<a>", value="null")],
                        blocks=[BlockFill(index=1, rows=[["Excel ?", "?", "x"]], answers=["None", "5"],
                                          paragraphs=["- Abrir o Excel\n- Salvar o arquivo", "Texto"])])
    normalize_fill(fill)
    assert fill.placeholders[0].value is None
    assert fill.blocks[0].rows == [["Excel", None, "x"]]
    assert fill.blocks[0].answers == [None, "5"]
    assert fill.blocks[0].paragraphs == ["Abrir o Excel", "Salvar o arquivo", "Texto"]


def test_analysis_drops_impossible_decisions(template):
    from template import validate_analysis
    analysis = analysis_for_fixture()
    analysis.blocks.append(BlockPlan(index=4, action="table_rows"))  # a paragraph can't get table rows
    analysis.blocks.append(BlockPlan(index=500, action="remove"))
    validate_analysis(template, analysis)
    assert all(b.index != 500 for b in analysis.blocks)
    assert not any(b.index == 4 and b.action == "table_rows" for b in analysis.blocks)
    # Only short placeholders found in kept text or headers remain.
    assert [p.text for p in analysis.placeholders] == ["<Cliente>", "<nome do processo>", "<Nome do processo>"]
    assert json.loads(analysis.model_dump_json())["step_format"]["label"] == "P{n}"
    assert len(body_blocks(Document(template))) == 18  # 17 blocks + section properties


def test_section_left_empty_by_the_analysis_gets_written(template):
    from template import validate_analysis
    analysis = analysis_for_fixture()
    for plan in analysis.blocks:
        if plan.index in (6, 7):  # "Sistemas Envolvidos": its instruction and its table both removed
            plan.action, plan.purpose = "remove", ""
    validate_analysis(template, analysis)
    plans = {p.index: p for p in analysis.blocks}
    assert plans[6].action == "write" and "Lista de todos os sistemas" in plans[6].purpose
    assert plans[15].action == "remove"  # the example steps' section is where the steps go, not "empty"
    assert plans[9].header_rows == 1 and plans[9].answer_column == 1
