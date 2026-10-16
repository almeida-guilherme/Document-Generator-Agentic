"""A small company-style PDD template, with the parts real templates have: placeholders (one split over
several runs, as Word often saves them), instructions, an example from another process, a question table."""
from docx import Document
from docx.shared import Inches, RGBColor
from PIL import Image


def build_template(path: str, image_path: str) -> str:
    Image.new("RGB", (400, 200), "#cccccc").save(image_path)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "<cliente> <Nome do processo>"
    doc.add_paragraph("<Cliente>", style="Title")                                             # 0
    doc.add_heading("Histórico de Revisão", level=1)                                         # 1
    table = doc.add_table(rows=2, cols=3)                                                    # 2
    for cell, text in zip(table.rows[0].cells, ["Versão", "Data", "Autor"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["0.1", "18/02/2022", "Maria Exemplo"]):
        cell.text = text
    doc.add_heading("Introdução", level=1)                                                   # 3
    intro = doc.add_paragraph("Este documento descreve o processo ")                         # 4
    for part in ("<nome ", "do processo", ">"):  # one placeholder split over three runs
        intro.add_run(part).font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    intro.add_run(" escolhido para automação.")
    doc.add_heading("Sistemas Envolvidos", level=1)                                          # 5
    doc.add_paragraph("<Lista de todos os sistemas envolvidos no processo>")                 # 6
    table = doc.add_table(rows=2, cols=3)                                                    # 7
    for cell, text in zip(table.rows[0].cells, ["#", "Aplicação", "Versão"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["1", "Outlook 2011", "2011"]):
        cell.text = text
    doc.add_heading("Métricas", level=1)                                                     # 8
    table = doc.add_table(rows=3, cols=2)                                                    # 9
    for row, texts in zip(table.rows, [["Métrica", "Descrição"], ["Quantas pessoas executam?", ""],
                                       ["Qual o risco de 0 a 10?", ""]]):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    doc.add_heading("Diagrama do Processo", level=1)                                         # 10
    doc.add_picture(image_path, width=Inches(3))                                             # 11
    doc.add_paragraph("Figura 1 - Exemplo de diagrama")                                      # 12
    doc.add_heading("Detalhamento do Processo", level=1)                                     # 13
    doc.add_heading("Login SAP", level=2)                                                    # 14
    step = doc.add_paragraph()                                                               # 15
    step.add_run("P1").bold = True
    step.add_run(" - Realizar abertura do SAP 760")
    doc.add_picture(image_path, width=Inches(3))                                             # 16
    doc.save(path)
    return path
