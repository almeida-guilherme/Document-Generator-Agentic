"""Offline tests: the LLM and Whisper are replaced by fakes, so no API key is needed."""
import json
import os
import re
import threading
import subprocess
from types import SimpleNamespace
import pytest
from docx import Document

from pipeline import run_pdd_pipeline
from schemas import (FrameCaption, ProcessOutline, OutlineStep, WindowRefinement, RefinedStep,
                     RequirementItem, ExceptionItem, KeyName, ProcessInput)
from transcribe import transcribe_audio, has_audio
from tests.synthetic_video import build


class FakeLLM:
    """Mimics `llm.with_structured_output(schema).invoke(messages)` with deterministic answers."""

    model_name = "fake"

    def __init__(self, fail_after_captions: int | None = None):
        self.refine_calls = 0
        self.captions = 0
        self.fail_after_captions = fail_after_captions
        self.lock = threading.Lock()

    def with_structured_output(self, schema):
        return SimpleNamespace(invoke=lambda messages: self._answer(schema, messages))

    def _answer(self, schema, messages):
        text = messages[-1].content
        if schema is FrameCaption:
            with self.lock:  # captions are requested from several threads
                if self.captions == self.fail_after_captions:
                    raise ConnectionError("simulated network failure")
                self.captions += 1
            return FrameCaption(screen_or_system="Intranet", observed_action="something", visible_text=[])
        if schema is ProcessOutline:
            ids = re.findall(r'"id": "(F\d+)"', text)
            times = dict(re.findall(r'"id": "(F\d+)", "t": ([\d.]+)', text))
            # One step per pair of frames (as a range), one range ending at an unknown id that the code must
            # ignore, and two steps with no frames: one never narrated (stays without image) and one narrated
            # just before the next step's frames.
            steps = [OutlineStep(instruction=f"You need to do {a}+{b}", short_label="Do it", first_frame_id=a,
                                 last_frame_id=b) for a, b in zip(ids[::2], ids[1::2])]
            steps[1].last_frame_id = "F999"
            steps.insert(1, OutlineStep(instruction="You need to wait for approval", short_label="Wait"))
            steps.insert(2, OutlineStep(instruction="You need to clear the form", short_label="Clear",
                                        narrated_at=float(times[ids[2]])))
            steps[0].application = steps[3].application = "Intranet"
            steps[3].decision = "If the client exists, open it; otherwise create it."
            steps[4].instruction += " in Report 5.3.26"  # misread name, fixed through key_names
            return ProcessOutline(
                process_name="Create invoice", project_proposal="Proposal.", steps=steps,
                key_names=[KeyName(name="Report 5.5.26", variants=["Report 5.3.26", "Report"])],
                inputs=[ProcessInput(name="Client list", kind="structured", source="Intranet")],
                requirements=[RequirementItem(category="frequency", description="Runs daily.")],
                business_exceptions=[ExceptionItem(name="No client", action="Client missing", action_to_be_taken="Skip"),
                                     ExceptionItem(name="no client ", action="dup", action_to_be_taken="dup")],
            )
        if schema is WindowRefinement:
            self.refine_calls += 1
            numbered = [int(n) for n in re.findall(r"^(\d+)\. ", text, re.M)]
            guesses = re.findall(r"first guess of frames: \[(.*?)\]", text)
            refined = []
            for n, guess in zip(numbered, guesses):
                ids = re.findall(r"F\d+", guess)
                if n == 5:
                    assert "- Report 5.5.26" in text  # the refinement is told the correct spellings
                    refined.append(RefinedStep(number=n, instruction="Refined 5 in Report 5.3.26", frame_ids=ids,
                                               main_frame_id=ids[-1]))
                    continue
                if n == 4:
                    continue  # forgotten by the model: code must keep the outline version
                main = "F998" if n == numbered[-1] else (ids[-1] if ids else "")  # last one invalid
                refined.append(RefinedStep(number=n, instruction=f"Refined {n}", frame_ids=ids, main_frame_id=main))
            return WindowRefinement(steps=refined)
        raise AssertionError(f"unexpected schema {schema}")


@pytest.fixture(scope="module")
def video(tmp_path_factory):
    path = str(tmp_path_factory.mktemp("video") / "synthetic.mp4")
    build(path)
    return path


def test_pipeline_end_to_end(video, tmp_path):
    llm = FakeLLM()
    output, report = run_pdd_pipeline(video, str(tmp_path), llm=llm, structure_llm=llm,
                                      transcribe=lambda v, audio_dir: [], max_frames_per_window=5)

    assert report["frames_kept"] >= 10
    assert report["steps"] == report["frames_kept"] // 2 + 2  # one per pair of frames + the two frameless steps
    assert llm.refine_calls == report["windows"] > 1
    assert report["steps_not_refined"] == [4]
    assert report["steps_framed_by_narration"] == [3]
    assert report["steps_without_image"] == [2]  # the step that was never narrated
    assert {"F999", "F998"} <= set(report["invalid_frame_refs"])
    assert report["images_inserted"] == report["steps"] - 1
    assert report["requirements"] == 1 and report["exceptions"] == 1  # duplicate exception removed

    assert report["names_fixed"] == {"Report 5.3.26": "Report 5.5.26"}
    assert report["applications"] == ["Intranet"] and report["decision_points"] == [4]
    assert report["recorded_seconds"] > 0

    doc = Document(output)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Decision: If the client exists" in text and "Step 1 — Intranet" in text
    overview = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    assert overview["Applications"] == "Intranet"
    assert overview["Inputs"] == "Client list (structured, from Intranet)"
    assert overview["Decision points"] == "step 4"
    assert len(doc.inline_shapes) == report["images_inserted"] + int(report["flowchart"])
    # Word refuses files where a paragraph's properties (w:pPr) aren't its first child.
    for p in doc.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        if p.pPr is not None:
            assert p[0] is p.pPr
    for name in ("captions.json", "transcript.json", "process_model.json", "frames.json", "report.json"):
        assert os.path.exists(tmp_path / name)


def test_steps_are_numbered_and_timed_by_code(video, tmp_path):
    llm = FakeLLM()
    run_pdd_pipeline(video, str(tmp_path), llm=llm, structure_llm=llm, transcribe=lambda v, audio_dir: [],
                     max_frames_per_window=5)
    steps = json.loads((tmp_path / "process_model.json").read_text())["as_is"]
    assert [s["number"] for s in steps] == list(range(1, len(steps) + 1))
    timed = [s["time"] for s in steps if s["frame_ids"]]
    assert timed == sorted(timed)
    assert steps[1]["time"] == "" and steps[1]["main_frame_path"] is None
    assert steps[2]["frame_ids"]  # found from its narration time
    assert steps[3]["instruction"].startswith("You need to do")  # not refined: outline text kept
    assert steps[4]["instruction"] == "Refined 5 in Report 5.5.26"  # misreading corrected
    # The screenshot is the frame the model chose, or the step's last frame when its choice is invalid.
    for s in steps:
        if s["main_frame_path"]:
            assert os.path.basename(s["main_frame_path"]) == f'{s["frame_ids"][-1]}.jpg'


def test_rerun_from_previous_run_reuses_captions(video, tmp_path):
    llm = FakeLLM()
    first = tmp_path / "first"; first.mkdir()
    run_pdd_pipeline(video, str(first), llm=llm, structure_llm=llm, transcribe=lambda v, audio_dir: [])
    second = tmp_path / "second"; second.mkdir()
    no_captions = SimpleNamespace(model_name="unused", with_structured_output=lambda schema: pytest.fail("captioned again"))
    _, report = run_pdd_pipeline(video, str(second), llm=no_captions, structure_llm=llm,
                                 transcribe=lambda v, audio_dir: pytest.fail("transcribed again"), from_run=str(first))
    assert report["images_inserted"] == report["steps"] - 1
    assert not (second / "frames").exists()


def test_transcription_offsets_chunks(tmp_path):
    media = str(tmp_path / "with_audio.mp4")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=white:s=320x240:d=25",
                    "-f", "lavfi", "-i", "sine=frequency=440:duration=25", "-shortest", media], check=True)
    assert has_audio(media)

    calls = []
    def create(**kwargs):
        calls.append(kwargs["file"].name)
        return SimpleNamespace(segments=[SimpleNamespace(text=f" chunk {len(calls)} ", start=1.0, end=2.0)])
    client = SimpleNamespace(audio=SimpleNamespace(transcriptions=SimpleNamespace(create=create)))

    segments = transcribe_audio(media, audio_dir=str(tmp_path / "audio"), client=client, chunk_seconds=10)
    assert len(calls) == 3
    assert [s["start"] for s in segments] == [1.0, 11.0, 21.0]
    assert segments[0]["text"] == "chunk 1"


def test_video_without_audio_skips_transcription(video, tmp_path):
    assert not has_audio(video)
    assert transcribe_audio(video, audio_dir=str(tmp_path), client=object()) == []


def test_out_of_order_frames_are_dropped():
    from schemas import Step
    from structure import enforce_chronology
    frames = {f"F{i:03d}": {"id": f"F{i:03d}", "timestamp": float(i * 10), "path": f"F{i:03d}.jpg"} for i in range(1, 31)}
    step = lambda n, ids, main: Step(number=n, time="", instruction="x", short_label="x", frame_ids=ids,
                                     main_frame_path=f"{main}.jpg")
    steps = [step(1, ["F001", "F002"], "F002"),
             step(2, ["F004", "F029"], "F029"),   # one frame from far later: dropped, screenshot falls back to F004
             step(3, ["F028"], "F028"),           # only a far-later frame: left without screenshot
             step(4, ["F006", "F007"], "F007")]
    dropped = enforce_chronology(steps, frames)
    assert dropped == {2: ["F029"], 3: ["F028"]}
    assert steps[1].frame_ids == ["F004"] and steps[1].main_frame_path == "F004.jpg" and steps[1].time == "00:40"
    assert steps[2].frame_ids == [] and steps[2].main_frame_path is None and steps[2].time == ""
    assert steps[3].main_frame_path == "F007.jpg"


def test_interrupted_run_resumes_where_it_stopped(video, tmp_path):
    failing = FakeLLM(fail_after_captions=4)
    with pytest.raises(ConnectionError):
        run_pdd_pipeline(video, str(tmp_path), llm=failing, structure_llm=failing, transcribe=lambda v, audio_dir: [])
    # Frames are captioned several at a time, but what is saved is always the frames in order up to the
    # first one that failed (which may be the very first frame, as the failing call can be any of them).
    captions_file = tmp_path / "captions.json"
    saved = [c["id"] for c in json.loads(captions_file.read_text())] if captions_file.exists() else []
    assert len(saved) <= 4 and saved == [f"F{i:03d}" for i in range(1, len(saved) + 1)]

    llm = FakeLLM()
    _, report = run_pdd_pipeline(video, str(tmp_path), llm=llm, structure_llm=llm,
                                 transcribe=lambda v, audio_dir: [])
    assert llm.captions == report["frames_kept"] - len(saved)  # only the frames not captioned before
    captions = json.loads((tmp_path / "captions.json").read_text())
    assert [c["id"] for c in captions] == [f"F{i:03d}" for i in range(1, report["frames_kept"] + 1)]

    # Resuming a finished run redoes nothing but the document.
    again = SimpleNamespace(model_name="unused", with_structured_output=lambda schema: pytest.fail("called the model"))
    output, _ = run_pdd_pipeline(video, str(tmp_path), llm=again, structure_llm=again,
                                 transcribe=lambda v, audio_dir: pytest.fail("transcribed again"))
    assert os.path.exists(output)


def test_durations_come_from_frame_times():
    from schemas import Step
    from structure import measure_durations
    frames = {f"F{i:03d}": {"timestamp": float(i * 10)} for i in range(1, 10)}
    step = lambda n, ids: Step(number=n, time="", instruction="x", short_label="x", frame_ids=ids)
    steps = [step(1, ["F001", "F002"]), step(2, []), step(3, ["F004", "F006"])]
    measure_durations(steps, frames)
    assert (steps[0].start_seconds, steps[0].duration_seconds) == (10.0, 30.0)  # until step 3 starts
    assert steps[1].start_seconds is None and steps[1].duration_seconds is None
    assert (steps[2].start_seconds, steps[2].duration_seconds) == (40.0, 20.0)  # last step: until its last frame


def test_costs_use_own_price_table():
    from costs import usage_cost, price_for
    assert price_for("gpt-4.1-mini-2025-04-14") == price_for("gpt-4.1-mini") != price_for("gpt-4.1")
    assert price_for("gpt-4.1-2025-04-14") == price_for("gpt-4.1")
    cost, unknown = usage_cost({
        "gpt-4.1-2025-04-14": {"input_tokens": 1_000_000, "output_tokens": 100_000,
                               "input_token_details": {"cache_read": 500_000}},
        "some-new-model": {"input_tokens": 10, "output_tokens": 10},
    })
    assert cost == pytest.approx(0.5 * 2.00 + 0.5 * 0.50 + 0.1 * 8.00)
    assert unknown == ["some-new-model"]


def test_long_flowchart_uses_columns(tmp_path):
    from schemas import ProcessModel, Step
    from flowchart import build_flowchart, columns_for
    from PIL import Image
    assert [columns_for(n) for n in (1, 10, 11, 25, 100)] == [1, 1, 2, 3, 4]
    steps = [Step(number=i, time="", instruction="x", short_label=f"Step {i}", frame_ids=[]) for i in range(1, 26)]
    path = build_flowchart(ProcessModel(process_name="p", project_proposal="", as_is=steps), str(tmp_path / "chart"))
    with Image.open(path) as img:
        assert img.width > img.height / 2  # three columns, not one tall strip


def test_only_misreadings_are_replaced():
    from schemas import ProcessModel, Step
    from structure import apply_key_names
    model = ProcessModel(
        process_name="Report 5.3.26", project_proposal="Open My Views, then XRef Table 2.",
        as_is=[Step(number=1, time="", instruction="Save Report 5.3.26 in My Views", short_label="Save", frame_ids=[])],
        key_names=[KeyName(name="Report 5.5.26", variants=["Report 5.3.26"]),
                   KeyName(name="APN Blocked Orders View", variants=["My Views"]),  # a different thing
                   ])
    fixed = apply_key_names(model)
    assert "Report 5.3.26" in fixed and "My Views" not in fixed
    assert model.as_is[0].instruction == "Save Report 5.5.26 in My Views"
    model.project_proposal = "Open 'Report 5.5.26' or 'Report 5.3.26', then Report 5.3.26/Report 5.5.26."
    apply_key_names(model)
    assert model.project_proposal == "Open 'Report 5.5.26', then Report 5.5.26."


def test_misreading_check():
    from structure import is_misreading
    copy = "AP GTS Blocks Report - 5.5.26 - Afternoon.xlsx"
    assert is_misreading("AP GTS Blocks Report - 5.3.26 - Afternoon.xlsx", copy)  # digit misread
    assert is_misreading("AP GTS Blocks Report - 5.5.26 - Athenoon.xlsx", copy)
    assert is_misreading("APN Supply Chain Cockp", "APN Supply Chain Cockpit")      # cut off
    assert not is_misreading("AP GTS Blocks Report - 5.5.26.xlsx", copy)            # the template: another file
    assert not is_misreading("My Views", "APN Blocked Orders View")
    assert not is_misreading("SAP", "SAP BC")


def test_similar_names_are_settled_by_looking_again(tmp_path):
    from PIL import Image
    from schemas import ProcessModel, Step
    from verify_names import verify_names, NameCheck
    frames, captions = [], []
    for i in range(1, 5):
        path = str(tmp_path / f"F00{i}.jpg")
        Image.new("RGB", (1920, 1080), "white").save(path)
        frames.append({"id": f"F00{i}", "path": path, "timestamp": float(i)})
    # Captions read the title bar both ways; F004 shows a file list with two real, similar files.
    captions = [{"id": "F001", "visible_text": ["Report 5.5.26.xlsx"], "observed_action": ""},
                {"id": "F002", "visible_text": ["Report 5.3.26.xlsx"], "observed_action": ""},
                {"id": "F003", "visible_text": ["Table 1", "Table 2"], "observed_action": ""},
                {"id": "F004", "visible_text": ["Table 2"], "observed_action": ""}]
    on_screen = {"F001": ["Report 5.5.26.xlsx"], "F002": ["Report 5.5.26.xlsx"], "F003": ["Table 1", "Table 2"],
                 "F004": ["Table 2"]}

    class Vision:
        calls = []
        def with_structured_output(self, schema):
            assert schema is NameCheck
            return SimpleNamespace(invoke=self.invoke)
        def invoke(self, messages):
            content = messages[0].content
            assert len(content) == 1 + 9  # the prompt and 3x3 overlapping full-resolution tiles
            self.calls.append(content)
            # The fake can't see pixels, so it answers from the frame being checked (tracked by call order).
            fid = order[len(self.calls) - 1]
            return NameCheck(present=on_screen[fid])

    order = ["F001", "F002", "F003", "F004"]
    model = ProcessModel(process_name="p", project_proposal="Uses Report 5.3.26.xlsx and Table 2.",
                         as_is=[Step(number=1, time="", instruction="Open Report 5.3.26.xlsx", short_label="Open",
                                     frame_ids=[])],
                         key_names=[KeyName(name="Report 5.5.26.xlsx"), KeyName(name="Report 5.3.26.xlsx"),
                                    KeyName(name="Table 1"), KeyName(name="Table 2")])
    report = verify_names(model, captions, frames, Vision())
    assert report["names_fixed_by_vision"] == {"Report 5.3.26.xlsx": "Report 5.5.26.xlsx"}
    assert model.as_is[0].instruction == "Open Report 5.5.26.xlsx"
    assert model.project_proposal == "Uses Report 5.5.26.xlsx and Table 2."  # both tables are real
    assert [k.name for k in model.key_names] == ["Report 5.5.26.xlsx", "Table 1", "Table 2"]
    assert model.key_names[0].variants == ["Report 5.3.26.xlsx"]
