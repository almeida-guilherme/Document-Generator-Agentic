"""
Company PDD templates: reading a template, deciding what each part of it is for (analysis, done once per
template) and writing the content for one process into it (fill, done once per run).

A template is a .docx the company already uses: cover, revision history, fixed text, tables, instructions
in angle brackets ("<Lista de todos os sistemas>") and usually leftovers of the process it was last used
for (screenshots, steps, systems, names). The analysis gives every block of the body one action:

    keep            fixed text that applies to any process (kept as is; placeholders in it are filled)
    remove          instructions to the author and examples from other processes
    write           a paragraph replaced by text written for the new process
    table_rows      a table whose data rows are replaced (header kept)
    table_answers   a question/answer table whose answer column is filled
    steps           where the detailed steps with screenshots go
    diagram         where the process flowchart goes

and lists the placeholders ("<Cliente>") with what they stand for. The analysis is saved as JSON and
reused for every recording documented with that template; render_template.py writes the result.
"""
import json
import re
import sys
from datetime import date
from typing import Literal, cast
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.messages import HumanMessage, SystemMessage
from pydantic import BaseModel, Field

from schemas import ProcessModel, repair_model

MAX_CELL_CHARS = 150
MAX_PLACEHOLDER_CHARS = 80

# --- Reading ---


def body_blocks(doc) -> list:
    """The body's top-level elements (paragraphs, tables, ...) in order; analysis indexes refer to these."""
    return list(doc.element.body.iterchildren())


def is_toc(paragraph: Paragraph) -> bool:
    """A table of contents entry ("toc 1", "toc 2", ...), not its heading."""
    return paragraph.style is not None and re.fullmatch(r"toc \d", paragraph.style.name.lower()) is not None


def has_image(element) -> bool:
    return bool(element.findall(".//" + qn("w:drawing")) or element.findall(".//" + qn("w:pict")))


def describe_template(path: str) -> str:
    """The template as text for the model: one line per block with its index and style, tables row by row."""
    doc = Document(path)
    lines = []
    for i, el in enumerate(body_blocks(doc)):
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, doc)
            if is_toc(p):
                continue  # the table of contents is rebuilt from the headings
            image = " [IMAGE]" if has_image(el) else ""
            if not p.text.strip() and not image:
                continue
            style = p.style.name if p.style is not None else "Normal"
            lines.append(f"[{i}] paragraph style={style!r}{image}: {p.text.strip()}")
        elif tag == "tbl":
            table = Table(el, doc)
            lines.append(f"[{i}] table, {len(table.rows)} rows x {len(table.columns)} columns:")
            for r, row in enumerate(table.rows):
                cells = [c.text.strip().replace("\n", " / ")[:MAX_CELL_CHARS] for c in row.cells]
                lines.append(f"      row {r}: " + " | ".join(cells))
        elif tag == "sdt" and "".join(el.itertext()).strip():
            lines.append(f"[{i}] content control: {''.join(el.itertext()).strip()[:300]}")
    headers = []
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            headers += [p.text.strip() for p in part.paragraphs if p.text.strip()]
    if headers:
        lines.append("Page headers/footers: " + " | ".join(dict.fromkeys(headers)))
    return "\n".join(lines)


# --- Analysis (once per template) ---

class BlockPlan(BaseModel):
    index: int = Field(description="The block's index, as given in brackets")
    action: Literal["keep", "remove", "write", "table_rows", "table_answers", "steps", "diagram"]
    purpose: str = Field(default="", description="For write, table_rows, table_answers: what this block must contain "
                                                 "for a new process (e.g. 'Brief description of the process', "
                                                 "'One row per system used in the process'). Empty otherwise.")
    header_rows: int = Field(default=1, description="Tables: how many rows at the top are headers to keep")
    answer_column: int = Field(default=1, description="table_answers: 0-based column that receives the answers")
    reason: str = Field(default="", description="For remove: why (instruction to the author, example from another "
                                                "process, ...)")


class Placeholder(BaseModel):
    text: str = Field(description="The placeholder exactly as written, brackets included, e.g. '<Cliente>'")
    meaning: str = Field(description="What it stands for, e.g. 'client company name'")


class StepFormat(BaseModel):
    example_step_index: int | None = Field(description="Index of the paragraph showing how one detailed step is "
                                                       "written (e.g. 'P1 - Open SAP...'), or null if there is none")
    example_group_heading_index: int | None = Field(description="Index of the heading that groups example steps "
                                                                "(e.g. 'Login SAP'), or null if steps aren't grouped")
    label: str = Field(description="How steps are numbered, with {n} for the number, e.g. 'P{n}', 'Passo {n}', '{n}.'")
    separator: str = Field(description="What separates the label from the step text, e.g. ' - ', ': ', ' '")


class Labels(BaseModel):
    decision: str = Field(description="The word 'Decision' in the template's language")
    start: str = Field(description="The word 'Start' (flowchart start) in the template's language")
    end: str = Field(description="The word 'End' (flowchart end) in the template's language")


class TemplateAnalysis(BaseModel):
    language: str = Field(description="The template's language, e.g. 'Brazilian Portuguese'")
    missing_label: str = Field(description="Short marker, in the template's language, for information the author "
                                           "must still fill in, e.g. '[A preencher]'")
    placeholders: list[Placeholder] = Field(default_factory=list, max_length=40)
    blocks: list[BlockPlan] = Field(max_length=600, description="One entry per block listed, in order")
    step_format: StepFormat
    labels: Labels


ANALYSIS_PROMPT = """You analyze a company's Process Design Document (PDD) template for an RPA team. The template will be filled automatically for new processes, documented from screen recordings and their narration.

The template is given block by block, each with its index in brackets. Templates are usually a copy of a document written for another process, so besides the structure they contain:
- fixed text that applies to any process (e.g. "This document describes the business process chosen to be automated..."), generic lists (e.g. standard requirements for any automation), table headers, question columns;
- instructions to the author, usually in angle brackets or written as guidance ("<List of all systems involved>", "Expected that the process owner documents below all business exceptions...");
- placeholders for short values ("<Cliente>", "<Nome do Processo>", "<nome> - <email>");
- leftovers of the previous process: example steps and screenshots, example systems (e.g. "Outlook 2011", "SAP 740"), example rows, names of people, dates, file paths, example diagrams and their captions.

Give one entry per block:
- keep: fixed text that applies to any process, headings of the document's sections, the cover, the table of contents heading, generic lists. Short placeholders inside kept text are filled separately, so a fixed sentence with a placeholder in it is kept, not rewritten.
- remove: instructions to the author, explanations of what a section should contain, examples from another process (steps, screenshots, example sub-headings such as "Login SAP" inside the detailed steps), captions of removed examples, and leftover empty structure.
- write: a paragraph whose content must be written for each process (typically an instruction like "<Brief description of the process>" that stands where the content goes). Give its purpose. When a section has guidance paragraphs and no paragraph that clearly stands for the content, mark the last guidance paragraph as write (with the purpose) and the others as remove, so every section that needs content gets it. Captions of tables that are filled (e.g. "Table 1 - Example of input spreadsheet") are write, with the purpose of captioning that table for the new process.
- table_rows: a table with a header and example or empty rows, to be replaced by rows for the new process (systems, inputs, files, exceptions, revision history, impacted teams...). Give its purpose and header_rows.
- table_answers: a table of questions or items in one column whose other column is answered for each process (e.g. metrics questions). Give answer_column.
- steps: not used in blocks. Detailed steps are described by step_format instead; mark the example steps, their screenshots and their group headings as remove.
- diagram: the example process diagram image, to be replaced by the new process's flowchart. Its caption is write.

When a section's content goes in a table, the guidance paragraphs above it (e.g. "<List of all systems involved>" right above the systems table) are remove, not write, so the content isn't given twice; a neutral introductory sentence that suits any process may be kept.

Keep the document's section headings even if their content is all removed. Never remove the cover, logo, revision history heading or section headings. The revision history table is table_rows (its rows are examples).

step_format: point to the paragraph that shows how one detailed step is written and the heading that groups steps, and give the numbering label and separator they use.
placeholders: every distinct short placeholder (a value such as a name, not an instruction) inside kept text, the cover and page headers, as written (each spelling once, e.g. both "<Cliente>" and "<cliente>" if both appear).
"""


def analyze_template(path: str, llm) -> TemplateAnalysis:
    """Decides what each block of the template is for. `llm` is a LangChain chat model."""
    analysis = repair_model(cast(TemplateAnalysis, llm.with_structured_output(TemplateAnalysis).invoke([
        SystemMessage(content=ANALYSIS_PROMPT),
        HumanMessage(content=f"Template:\n\n{describe_template(path)}"),
    ])))
    validate_analysis(path, analysis)
    return analysis


def validate_analysis(path: str, analysis: TemplateAnalysis) -> None:
    """Drops decisions that can't apply (unknown indexes, table actions on paragraphs and vice versa)."""
    doc = Document(path)
    blocks = body_blocks(doc)
    valid = []
    for plan in analysis.blocks:
        if not 0 <= plan.index < len(blocks):
            continue
        tag = blocks[plan.index].tag.split("}")[1]
        if plan.action.startswith("table_") and tag != "tbl":
            continue
        if plan.action in ("write", "diagram", "steps") and tag != "p":
            continue
        if plan.action == "steps":
            plan.action = "remove"
        if plan.action.startswith("table_"):
            columns = len(Table(blocks[plan.index], doc).columns)
            plan.header_rows = max(1, plan.header_rows)  # replacing a table's header is never intended
            if plan.action == "table_answers" and not 0 < plan.answer_column < columns:
                plan.answer_column = columns - 1  # questions are in the first column, answers after them
        valid.append(plan)
    analysis.blocks = valid
    _fill_emptied_sections(doc, blocks, analysis)
    # Placeholders are only filled in kept text and page headers; the rest is rewritten or removed anyway.
    kept = [Paragraph(blocks[p.index], doc).text for p in valid if p.action == "keep" and blocks[p.index].tag == qn("w:p")]
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            kept += [p.text for p in part.paragraphs]
    kept_text = "\n".join(kept).lower()
    analysis.placeholders = [p for p in analysis.placeholders
                             if p.text.strip() and len(p.text) <= MAX_PLACEHOLDER_CHARS and p.text.lower() in kept_text]
    fmt = analysis.step_format
    for name in ("example_step_index", "example_group_heading_index"):
        index = getattr(fmt, name)
        if index is not None and not (0 <= index < len(blocks) and blocks[index].tag == qn("w:p")):
            setattr(fmt, name, None)
    if "{n}" not in fmt.label:
        fmt.label = "{n}."


def _heading_level(doc, el) -> int | None:
    if el.tag != qn("w:p"):
        return None
    p = Paragraph(el, doc)
    name = p.style.name.lower() if p.style is not None else ""
    if name.startswith("heading") and p.text.strip():
        return int(name.split()[-1]) if name.split()[-1].isdigit() else 1
    return None


def _fill_emptied_sections(doc, blocks: list, analysis: TemplateAnalysis) -> None:
    """
    A section whose only content is an instruction to the author ("<lista de todos os arquivos de saída>")
    would come out empty if the instruction is removed, as the analysis sometimes decides. The last such
    paragraph is written for the process instead, with the instruction as its purpose.
    """
    plans = {p.index: p for p in analysis.blocks}
    fmt = analysis.step_format
    steps_at = {i for i in (fmt.example_step_index, fmt.example_group_heading_index) if i is not None}
    section: list[int] = []
    heading = None

    def close():
        if heading is None or plans.get(heading) is None or plans[heading].action != "keep" or steps_at & set(section):
            return
        if any(plans.get(i) is not None and plans[i].action != "remove" for i in section):
            return
        removed = [i for i in section if plans.get(i) is not None and plans[i].action == "remove"
                   and blocks[i].tag == qn("w:p") and Paragraph(blocks[i], doc).text.strip()]
        if removed:
            plan = plans[removed[-1]]
            plan.action, plan.purpose = "write", plan.purpose or Paragraph(blocks[plan.index], doc).text.strip()[:300]
            plan.reason = ""

    for i, el in enumerate(blocks):
        if _heading_level(doc, el) is not None:
            close()
            heading, section = i, []
        else:
            section.append(i)
    close()


def summarize_analysis(path: str, analysis: TemplateAnalysis) -> list[dict]:
    """Per block: its text, the action and why, for showing people what the analysis decided."""
    doc = Document(path)
    blocks = body_blocks(doc)
    rows = []
    for plan in analysis.blocks:
        el = blocks[plan.index]
        if el.tag == qn("w:tbl"):
            table = Table(el, doc)
            text = " | ".join(c.text.strip() for c in table.rows[0].cells) if table.rows else ""
            kind = "table"
        else:
            p = Paragraph(el, doc)
            text = p.text.strip() or ("[image]" if has_image(el) else "")
            kind = "heading" if p.style is not None and p.style.name.lower().startswith("heading") else "paragraph"
        rows.append({"index": plan.index, "kind": kind, "text": text[:200], "action": plan.action,
                     "purpose": plan.purpose, "reason": plan.reason})
    # Where the detailed steps go: in place of the template's example steps, or at the end.
    fmt = analysis.step_format
    anchors = [i for i in (fmt.example_group_heading_index, fmt.example_step_index) if i is not None]
    at = min(anchors) if anchors else len(blocks)
    label = fmt.label.replace("{n}", "1") + fmt.separator
    grouped = " grouped under headings like the example's" if fmt.example_group_heading_index is not None else ""
    position = next((n for n, row in enumerate(rows) if row["index"] >= at), len(rows))
    rows.insert(position, {"index": at, "kind": "steps", "text": f"{label}…", "action": "steps",
                           "purpose": f"The process steps, each written as \"{label}…\" with its screenshot{grouped}",
                           "reason": ""})
    return rows


# --- Fill (once per run) ---

class BlockFill(BaseModel):  # after normalize_fill, None marks values the author must fill in
    index: int = Field(description="Index of the block being filled")
    paragraphs: list[str | None] = Field(default_factory=list, max_length=40, description="write: the paragraphs replacing the block; '?' if the recording doesn't say")
    rows: list[list[str | None]] = Field(default_factory=list, max_length=80, description="table_rows: one list of cell values per row, "
                                  "one value per column; '?' where the recording doesn't say")
    answers: list[str | None] = Field(default_factory=list, max_length=80, description="table_answers: one answer per question row, in "
                               "order; '?' where the recording doesn't say")


class PlaceholderValue(BaseModel):
    text: str = Field(description="The placeholder exactly as given")
    value: str | None = Field(description="Its value for this process, or '?' if the recording doesn't say")


class StepGroup(BaseModel):
    title: str = Field(description="Short heading for a group of consecutive steps, e.g. 'Login SAP', 'Validate the email'")
    first_step: int = Field(description="Number of the first step of the group")


class TemplateFill(BaseModel):
    placeholders: list[PlaceholderValue] = Field(default_factory=list, max_length=40)
    blocks: list[BlockFill] = Field(default_factory=list, max_length=200)
    step_groups: list[StepGroup] = Field(default_factory=list, max_length=40, description="Groups of consecutive steps, in order; "
                                         "the first starts at step 1")


FILL_PROMPT = """You fill a company's Process Design Document (PDD) template for one process. You receive:
1. The process model: the steps (already written), applications, inputs, key names, requirements and exceptions, worked out from a recording of the process.
2. The narration of the recording (a meeting or a walkthrough), with timestamps.
3. The template parts to fill, each with its index, the section it is in, what it must contain and, for tables, the columns or the questions.
4. The placeholders to fill.

Write everything in {language}, in the template's tone (formal, impersonal, as in a process document).

Rules:
- Use only what the process model and the narration state or clearly show. When something isn't stated (client, contact names, versions, login modules, technology, times, risk scores...), write exactly ? (a single question mark) as that cell, answer, paragraph or placeholder value; never invent it and never copy the template's example value. Don't write phrases like "not informed" or "não informado" either: use ?, and the author is shown what is left to complete.
- write: give the paragraphs, one string per paragraph; for a list, one string per item, without bullet characters.
- Captions: describe the filled table or the diagram for this process, numbered as in the template (e.g. "Tabela 1 - Campos da planilha de entrada"). Never keep the words "Exemplo"/"Example" from the template's caption.
- table_rows: one value per column, in the order of the columns given. Each value has the kind and length of the template's example value for that column: where the example is a short level or code (e.g. impact "Médio", required "SIM", "c01"), write a short level or code of that kind (e.g. "Alto", "Médio", "Baixo"), not a sentence; descriptions go in the columns meant for them. Fill the rows the process needs (e.g. one row per application, per input file, per exception). Revision history: one row for the first version, dated {today}, describing the document's creation from the recording; null for the author.
- table_answers: one answer per question, in order. Volumes, frequencies, times and people come from the narration when they are said.
- Exceptions tables: use the process model's exceptions plus situations the narration describes where the normal flow doesn't apply or something can fail, with the step they happen in (e.g. "P4") when known.
- In/out of scope: what the narration says will or won't be automated; if it says nothing about scope, the activities of the steps are in scope.
- step_groups: split the steps into a few consecutive phases (usually by application or goal, 2 to 8 steps each), with short titles in the style of the template's example group headings.
"""


def _section_paths(doc, blocks: list) -> dict[int, str]:
    """For each block index, the headings it is under ("Visão geral > Entradas")."""
    path: list[tuple[int, str]] = []
    result = {}
    for i, el in enumerate(blocks):
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            name = p.style.name.lower() if p.style is not None else ""
            if name.startswith("heading") and p.text.strip():
                level = int(name.split()[-1]) if name.split()[-1].isdigit() else 1
                path = [h for h in path if h[0] < level] + [(level, p.text.strip())]
        result[i] = " > ".join(h[1] for h in path)
    return result


def describe_fill_targets(path: str, analysis: TemplateAnalysis) -> str:
    doc = Document(path)
    blocks = body_blocks(doc)
    sections = _section_paths(doc, blocks)
    lines = []
    for plan in analysis.blocks:
        if plan.action not in ("write", "table_rows", "table_answers"):
            continue
        el = blocks[plan.index]
        lines.append(f"[{plan.index}] {plan.action} in section '{sections[plan.index]}': {plan.purpose}")
        if plan.action == "write":
            lines.append(f"    template text: {Paragraph(el, doc).text.strip()[:300]}")
        else:
            table = Table(el, doc)
            header = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
            if plan.action == "table_rows":
                lines.append(f"    columns: {header}")
                example = [c.text.strip()[:80] for c in table.rows[plan.header_rows].cells] \
                    if len(table.rows) > plan.header_rows else []
                if any(example):
                    # Shows the kind and length of value each column expects (e.g. "Médio" for impact).
                    lines.append(f"    the template's example row, from another process (match the kind and length "
                                 f"of each value, don't copy it): {example}")
            else:
                questions = [" | ".join(c.text.strip()[:200] for i, c in enumerate(row.cells) if i != plan.answer_column)
                             for row in table.rows[plan.header_rows:]]
                lines.append("    questions:\n" + "\n".join(f"      {n + 1}. {q}" for n, q in enumerate(questions)))
    return "\n".join(lines)


def _model_text(model: ProcessModel) -> str:
    data = model.model_dump(exclude={"key_names"})
    for step in data["as_is"]:
        for key in ("frame_ids", "main_frame_path", "start_seconds", "duration_seconds", "short_label"):
            step.pop(key, None)
    data["key_names"] = [k.name for k in model.key_names]
    return json.dumps(data, ensure_ascii=False, indent=1)


def fill_template(path: str, analysis: TemplateAnalysis, model: ProcessModel, transcript: list[dict],
                  llm, today: date | None = None) -> TemplateFill:
    """Writes the content of every part of the template the analysis marked to fill."""
    narration = "\n".join(f"[{t['start']:.0f}s] {t['text']}" for t in transcript) or "(no narration)"
    placeholders = "\n".join(f"- {p.text}: {p.meaning}" for p in analysis.placeholders) or "(none)"
    fill = repair_model(cast(TemplateFill, llm.with_structured_output(TemplateFill).invoke([
        SystemMessage(content=FILL_PROMPT.format(language=analysis.language,
                                                 today=(today or date.today()).strftime("%d/%m/%Y"))),
        HumanMessage(content=f"Process model:\n{_model_text(model)}\n\nNarration:\n{narration}\n\n"
                             f"Template parts to fill:\n{describe_fill_targets(path, analysis)}\n\n"
                             f"Placeholders:\n{placeholders}"),
    ])))
    normalize_fill(fill)
    steps = len(model.as_is)
    fill.step_groups = sorted({g.first_step: g for g in fill.step_groups if 1 <= g.first_step <= steps}.values(),
                              key=lambda g: g.first_step)
    return fill


UNKNOWN = {"?", "null", "none", "", "n/d"}


def _known(value: str | None) -> str | None:
    """None for the model's "don't know" answers ('?', or 'null' written as text); "Excel ?" becomes "Excel"."""
    if value is None or value.strip().lower() in UNKNOWN:
        return None
    # "Excel ?" (unknown version) becomes "Excel"; "Ana -" (unknown email after "<nome> - <email>") "Ana".
    return re.sub(r"(\s+\?|\s*[-–:,])+$", "", value.strip()) or None


def normalize_fill(fill: TemplateFill) -> None:
    """Unknown values become None; paragraphs written as one bulleted string become one paragraph per item."""
    for p in fill.placeholders:
        p.value = _known(p.value)
    for block in fill.blocks:
        block.rows = [[_known(v) for v in row] for row in block.rows]
        block.answers = [_known(a) for a in block.answers]
        paragraphs = []
        for text in block.paragraphs:
            for line in (text or "").split("\n"):
                line = re.sub(r"^\s*(?:[-•*]|\d+[.)])\s+", "", line) if "\n" in text else line
                if line.strip():
                    paragraphs.append(_known(line))
        block.paragraphs = paragraphs


if __name__ == "__main__":
    # python template.py TEMPLATE.docx: analyzes a template and prints what will be kept, removed and filled.
    from dotenv import load_dotenv
    from langchain_openai import ChatOpenAI
    load_dotenv()
    template_path = sys.argv[1]
    result = analyze_template(template_path, ChatOpenAI(model="gpt-4.1", temperature=0))
    print(f"Language: {result.language} | missing marker: {result.missing_label}")
    print("Placeholders:", ", ".join(f"{p.text} ({p.meaning})" for p in result.placeholders))
    print("Steps:", result.step_format.model_dump())
    for row in summarize_analysis(template_path, result):
        note = row["purpose"] or row["reason"]
        print(f"[{row['index']:3}] {row['action']:13} {row['kind']:9} {row['text'][:70]!r}  {note[:80]}")
    out = template_path + ".analysis.json"
    with open(out, "w", encoding="utf-8") as f:
        json.dump(result.model_dump(), f, ensure_ascii=False, indent=2)
    print("Saved", out)
