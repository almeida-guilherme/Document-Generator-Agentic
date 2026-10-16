"""
Company PDD templates: uploaded once, analyzed once (what to keep, remove and fill; see template.py), then
chosen when uploading recordings. Each run gets its own copy of the template and its analysis, so deleting
or replacing a template never changes documents already made with it.
"""
import json
import os
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from pydantic import BaseModel
from sqlmodel import select

from web.auth import current_user
from web.db import Template, User, session
from web.keys import key_value, pick_key

router = APIRouter(prefix="/api/templates", tags=["templates"])

TEMPLATE_FILE = "template.docx"
ANALYSIS_FILE = "template_analysis.json"
MAX_TEMPLATE_MB = 25


class TemplateOut(BaseModel):
    id: str
    name: str
    language: str
    created_at: datetime


class TemplateBlock(BaseModel):
    index: int
    kind: str
    text: str
    action: str
    purpose: str
    reason: str


class TemplateDetail(TemplateOut):
    missing_label: str
    placeholders: list[dict]
    step_label: str
    blocks: list[TemplateBlock]


def default_analyze(path: str, api_key: str | None = None):
    """Analyzes a template with the user's API key (None: the server's)."""
    from langchain_openai import ChatOpenAI
    from pipeline import STRUCTURE_MODEL
    from template import analyze_template
    return analyze_template(path, ChatOpenAI(model=STRUCTURE_MODEL, temperature=0, max_retries=5, api_key=api_key))


def template_dir(request: Request, template: Template) -> Path:
    return request.app.state.settings.templates_dir / template.id


def own_template(request: Request, template_id: str, user: User) -> Template:
    with session(request.app.state.engine) as db:
        template = db.get(Template, template_id)
    if template is None or template.user_id != user.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Template not found.")
    return template


def copy_into_run(request: Request, template: Template, run_dir: Path) -> None:
    """The run's own copy, which the pipeline finds in the run folder (pipeline.TEMPLATE_FILE)."""
    from pipeline import TEMPLATE_ANALYSIS_FILE, TEMPLATE_FILE as RUN_TEMPLATE_FILE
    source = template_dir(request, template)
    shutil.copyfile(source / TEMPLATE_FILE, run_dir / RUN_TEMPLATE_FILE)
    shutil.copyfile(source / ANALYSIS_FILE, run_dir / TEMPLATE_ANALYSIS_FILE)


@router.get("", response_model=list[TemplateOut])
def list_templates(request: Request, user: User = Depends(current_user)):
    with session(request.app.state.engine) as db:
        return db.exec(select(Template).where(Template.user_id == user.id).order_by(Template.created_at.desc())).all()


@router.post("", status_code=status.HTTP_201_CREATED, response_model=TemplateDetail)
def create_template(request: Request, file: UploadFile = File(...), user: User = Depends(current_user)):
    name = os.path.basename(file.filename or "template.docx")
    if not name.lower().endswith(".docx"):
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Upload the template as a Word document (.docx).")
    api_key = key_value(request, pick_key(request, user, None))  # the user's newest key, else the server's
    template = Template(user_id=user.id, name=name, language="")
    folder = template_dir(request, template)
    folder.mkdir(parents=True)
    try:
        data = file.file.read(MAX_TEMPLATE_MB * 1024 * 1024 + 1)
        if len(data) > MAX_TEMPLATE_MB * 1024 * 1024:
            raise HTTPException(status.HTTP_413_CONTENT_TOO_LARGE, f"The template is larger than {MAX_TEMPLATE_MB} MB.")
        path = folder / TEMPLATE_FILE
        path.write_bytes(data)
        if not zipfile.is_zipfile(path):
            raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "The file could not be read as a Word document.")
        try:
            from docx import Document
            Document(str(path))
        except Exception:
            raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "The file could not be read as a Word document.")
        analysis = request.app.state.analyze_template(str(path), api_key)
        (folder / ANALYSIS_FILE).write_text(json.dumps(analysis.model_dump(), ensure_ascii=False, indent=2),
                                            encoding="utf-8")
    except BaseException:
        shutil.rmtree(folder, ignore_errors=True)
        raise
    template.language = analysis.language
    with session(request.app.state.engine) as db:
        db.add(template)
        db.commit()
    return _detail(request, template)


def _detail(request: Request, template: Template) -> TemplateDetail:
    from template import TemplateAnalysis, summarize_analysis
    folder = template_dir(request, template)
    analysis = TemplateAnalysis.model_validate_json((folder / ANALYSIS_FILE).read_text(encoding="utf-8"))
    blocks = summarize_analysis(str(folder / TEMPLATE_FILE), analysis)
    return TemplateDetail(**template.model_dump(include=set(TemplateOut.model_fields)),
                          missing_label=analysis.missing_label,
                          placeholders=[p.model_dump() for p in analysis.placeholders],
                          step_label=analysis.step_format.label.replace("{n}", "1") + analysis.step_format.separator,
                          blocks=[TemplateBlock(**b) for b in blocks])


@router.get("/{template_id}", response_model=TemplateDetail)
def get_template(template_id: str, request: Request, user: User = Depends(current_user)):
    return _detail(request, own_template(request, template_id, user))


@router.delete("/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_template(template_id: str, request: Request, user: User = Depends(current_user)):
    template = own_template(request, template_id, user)
    with session(request.app.state.engine) as db:
        db.delete(db.get(Template, template_id))
        db.commit()
    shutil.rmtree(template_dir(request, template), ignore_errors=True)
